"""
convert_markdown_to_docs.py
High-fidelity Markdown to DOCX and PDF converter using python-docx and Microsoft Word COM.
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client


def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Set clean subtle borders on a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_formatted_text(p, text):
    """Parse inline markdown formatting (bold, italic, code) and append runs to paragraph."""
    # Pattern to match **bold**, `code`, *italic*
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(14, 116, 144)  # Cyan/teal accent
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        else:
            p.add_run(token)


def markdown_to_docx(md_path, docx_path, doc_title="SIH26170 Documentation"):
    """Convert Markdown document into a styled Word (.docx) document."""
    doc = Document()

    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SIH26170 | AI-Driven Semiconductor Burn-in Screening")
        hrun.font.name = 'Segoe UI'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun = fp.add_run("Smart India Hackathon 2026 — Problem Statement 26170 (ISRO)")
        frun.font.name = 'Segoe UI'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)  # Slate-800
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        # Parse table rows
        rows_data = []
        for line in table_lines:
            # Strip outer pipes
            cleaned = line.strip()
            if cleaned.startswith('|'):
                cleaned = cleaned[1:]
            if cleaned.endswith('|'):
                cleaned = cleaned[:-1]
            cells = [c.strip() for c in cleaned.split('|')]
            # Skip separator line like |---|---|
            if cells and all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            rows_data.append(cells)

        if len(rows_data) > 0:
            num_cols = max(len(r) for r in rows_data)
            table = doc.add_table(rows=len(rows_data), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table)

            for row_idx, row in enumerate(rows_data):
                is_header = (row_idx == 0)
                for col_idx in range(num_cols):
                    cell = table.cell(row_idx, col_idx)
                    val = row[col_idx] if col_idx < len(row) else ""
                    cell.text = ""
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

                    if is_header:
                        set_cell_background(cell, "1E293B")  # Dark slate
                        add_formatted_text(cp, val)
                        for r in cp.runs:
                            r.bold = True
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg_color)
                        add_formatted_text(cp, val)
                        for r in cp.runs:
                            r.font.size = Pt(9)
                            if not r.font.color.rgb:
                                r.font.color.rgb = RGBColor(51, 65, 85)

            # Add spacing after table
            sp_p = doc.add_paragraph()
            sp_p.paragraph_format.space_before = Pt(0)
            sp_p.paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    def flush_code_block():
        nonlocal in_code_block, code_lines
        if not code_lines:
            return
        code_text = "".join(code_lines)
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(6)
        cp.paragraph_format.left_indent = Inches(0.25)
        cp.paragraph_format.right_indent = Inches(0.25)

        run = cp.add_run(code_text.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(15, 23, 42)

        # Style with light shading
        pPr = cp._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="18" w:space="8" w:color="0284C7"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        pPr.append(shd)

        code_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if '|' in stripped and not stripped.startswith('//'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal Rule
        if stripped in ('---', '***', '___'):
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(6)
            hp.paragraph_format.space_after = Pt(8)
            pPr = hp._p.get_or_add_pPr()
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="E2E8F0"/>'
                f'</w:pBdr>'
            )
            pPr.append(pBdr)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.name = 'Segoe UI Semibold'
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(15, 23, 42)  # #0F172A
            i += 1
            continue

        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:])
            run.bold = True
            run.font.name = 'Segoe UI Semibold'
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(30, 41, 59)  # #1E293B
            i += 1
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.name = 'Segoe UI Semibold'
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(51, 65, 85)  # #334155
            i += 1
            continue

        if stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.name = 'Segoe UI'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Blockquote
        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.2)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="18" w:space="8" w:color="0284C7"/>'
                f'</w:pBdr>'
            )
            pPr.append(pBdr)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
            pPr.append(shd)

            text = stripped[2:]
            add_formatted_text(p, text)
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Bullet List (- or *)
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text = stripped[2:]
            add_formatted_text(p, text)
            i += 1
            continue

        # Numbered List (1. 2. etc)
        match_num = re.match(r'^\d+\.\s+(.*)$', stripped)
        if match_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, match_num.group(1))
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_text(p, stripped)
        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path} ({os.path.getsize(docx_path)} bytes)")


def convert_docx_to_pdf(docx_path, pdf_path):
    """Convert .docx file to .pdf using native Microsoft Word COM automation."""
    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)

    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(abs_docx)
        # 17 corresponds to wdFormatPDF
        doc.SaveAs(abs_pdf, FileFormat=17)
        doc.Close()
        print(f"Generated PDF:  {pdf_path} ({os.path.getsize(abs_pdf)} bytes)")
    except Exception as e:
        print(f"Error converting {docx_path} to PDF: {e}")
    finally:
        word.Quit()


def main():
    docs_dir = r"C:\Users\admin\Desktop\sih26170\docs"

    files_to_convert = [
        {
            "md": os.path.join(docs_dir, "SIH26170_Prototype_Documentation.md"),
            "docx": os.path.join(docs_dir, "SIH26170_Prototype_Documentation.docx"),
            "pdf": os.path.join(docs_dir, "SIH26170_Prototype_Documentation.pdf"),
            "title": "SIH26170 Prototype & System Architecture Documentation"
        },
        {
            "md": os.path.join(docs_dir, "SIH26170_File_and_Data_Documentation.md"),
            "docx": os.path.join(docs_dir, "SIH26170_File_and_Data_Documentation.docx"),
            "pdf": os.path.join(docs_dir, "SIH26170_File_and_Data_Documentation.pdf"),
            "title": "SIH26170 File Structure & Data Documentation"
        },
        {
            "md": os.path.join(docs_dir, "SIH26170_Deployment_and_Architecture_Guide.md"),
            "docx": os.path.join(docs_dir, "SIH26170_Deployment_and_Architecture_Guide.docx"),
            "pdf": os.path.join(docs_dir, "SIH26170_Deployment_and_Architecture_Guide.pdf"),
            "title": "SIH26170 Deployment & Architecture Guide"
        }
    ]

    for item in files_to_convert:
        print(f"\nProcessing: {os.path.basename(item['md'])}")
        markdown_to_docx(item["md"], item["docx"], item["title"])
        convert_docx_to_pdf(item["docx"], item["pdf"])


if __name__ == "__main__":
    main()
